"""Lập bảng Excel đối chiếu phiếu chỉnh định rơ le với máy cắt / Recloser trên sơ đồ.

    node tools/chinh-dinh/doi-chieu.ts /tmp/doi-chieu.json
    python3 tools/chinh-dinh/xuat-excel.py /tmp/doi-chieu.json docs/doi-chieu-phieu-chinh-dinh.xlsx
"""
import json, re, sys, collections, datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

D = json.load(open(sys.argv[1], encoding="utf-8"))
RA = sys.argv[2]
TEN = D["tenTram"]

def khu(k):
    return "Bắc Kạn" if k.startswith("E26") else "Thái Nguyên"

def ten_tram(ma):
    t = TEN.get(ma, "")
    t = re.sub(r"^(SƠ ĐỒ\s*)?TRẠM\s*", "", t, flags=re.I)
    t = re.sub(r"\s*\(?[EA]\d+\.\d+\)?\s*$", "", t)
    return t.strip()

def thu_tu(k):
    m = re.match(r"E(\d+)\.(\d+):(\D*)(\d+)(.*)", k)
    if not m:
        return (999, 0, k)
    cot = re.match(r"/(\d+)", m.group(5) or "")
    return (int(m.group(1)), int(m.group(2)), int(m.group(4)), int(cot.group(1)) if cot else -1, k)

# phiếu chưa gặp thiết bị, theo lộ -> gợi ý khi số cột lệch
cung_lo = collections.defaultdict(list)
for p in D["phieuKhongTB"]:
    for k in filter(None, p["khoa"].split("; ")):
        if "/" in k:
            cung_lo[k.split("/")[0]].append(f'{p["tb"]} ({p["so"]})')

# thiết bị có khoá, chưa có phiếu (gộp trùng khoá)
gap = {}
for t in D["thietBi"]:
    if not t["khoa"] or t["soPhieu"]:
        continue
    g = gap.setdefault(t["khoa"], dict(t, soLan=0))
    g["soLan"] += 1
ds_tb = sorted(gap.values(), key=lambda t: thu_tu(t["khoa"]))
trung_ap = [t for t in ds_tb if t["kv"] <= 35]
cao_ap = [t for t in ds_tb if t["kv"] > 35]

wb = Workbook()
DAM = Font(bold=True)
NEN = PatternFill("solid", fgColor="DDEBF7")
KE = Border(*(Side(style="thin", color="999999"),) * 4)

def bang(ws, tieu_de, cot, dong, rong):
    ws.append([tieu_de])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Lập ngày {datetime.date.today():%d/%m/%Y} - đối chiếu {D['tongPhieu']} phiếu chỉnh định (Drive: Role/Phieu chinh dinh role) với sơ đồ kết dây trong phần mềm"])
    ws["A2"].font = Font(italic=True, size=9)
    ws.append(cot)
    for c in ws[3]:
        c.font = DAM; c.fill = NEN; c.border = KE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in dong:
        ws.append(r)
        for c in ws[ws.max_row]:
            c.border = KE
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for i, w in enumerate(rong, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(cot))}{ws.max_row}"

def dong_tb(ds):
    out = []
    for i, t in enumerate(ds, 1):
        tram, so = t["khoa"].split(":", 1)
        lo = so.split("/")[0]
        vi_tri = f'Trạm {tram} {ten_tram(tram)}'.strip() if t["tram"] else f'ĐZ {lo} {tram}'
        goi_y = "; ".join(cung_lo.get(f"{tram}:{lo}", [])) if "/" in so else ""
        out.append([i, khu(tram), vi_tri, t["loai"], f'{t["kv"]}kV', t["nhan"], f"MC {so} {tram}" if t["tram"] else t["nhan"].replace("MC ", "MC ", 1), t["to"], goi_y])
    return out

COT_TB = ["STT", "Khu vực", "Trạm / Đường dây", "Loại thiết bị", "Cấp điện áp", "Nhãn trên sơ đồ", "Thiết bị", "Tờ sơ đồ", "Phiếu cùng lộ nhưng khác số cột (cần kiểm tra)"]
RONG_TB = [6, 12, 30, 16, 10, 22, 24, 22, 60]

ws = wb.active
ws.title = "MC-REC trung áp chưa có phiếu"
bang(ws, "DANH SÁCH MÁY CẮT / RECLOSER TRUNG ÁP CHƯA TÌM THẤY PHIẾU CHỈNH ĐỊNH", COT_TB, dong_tb(trung_ap), RONG_TB)

ws = wb.create_sheet("MC 110-220kV chưa có phiếu")
bang(ws, "DANH SÁCH MÁY CẮT 110-220kV CHƯA TÌM THẤY PHIẾU CHỈNH ĐỊNH TRONG THƯ MỤC", COT_TB, dong_tb(cao_ap), RONG_TB)

def ly_do(p):
    tb = p["tb"]
    if " TG " in tb:
        return "Trạm trung gian chưa vẽ trên sơ đồ kết dây"
    if tb.startswith("TC ") or "sa thải" in tb:
        return "Bảo vệ thanh cái / tủ sa thải - không gắn với một máy cắt"
    if tb.startswith("Tụ bù"):
        return "Tụ bù trên đường dây - sơ đồ không vẽ máy cắt tụ"
    if not p["khoa"]:
        return "Không xác định được máy cắt từ tên ghi trên phiếu"
    return f'Không thấy {p["khoa"].replace(":", " - MC ")} trên sơ đồ (chưa vẽ / đã thay đổi số hiệu)'

ds_p = sorted(D["phieuKhongTB"], key=lambda p: (ly_do(p), p["tb"]))
ws = wb.create_sheet("Phiếu chưa có trên sơ đồ")
bang(ws, "PHIẾU CHỈNH ĐỊNH CHƯA GẮN ĐƯỢC VÀO THIẾT BỊ TRÊN SƠ ĐỒ",
     ["STT", "Thiết bị ghi trên phiếu", "Số phiếu", "Ngày BH", "Rơ le", "Tên file", "Lý do", "Link phiếu"],
     [[i, p["tb"], p["so"], p["nb"], p["rl"], p["ten"], ly_do(p), f'https://drive.google.com/file/d/{p["id"]}/view'] for i, p in enumerate(ds_p, 1)],
     [6, 32, 24, 11, 22, 34, 46, 40])

ds_l = sorted(D["luuY"], key=lambda p: p["tb"])
ws = wb.create_sheet("Phiếu cần lưu ý")
bang(ws, "PHIẾU CÓ SAI LỆCH / TRÍCH ĐOẠN CHƯA ĐỦ / BẢN LƯU CŨ",
     ["STT", "Thiết bị", "Số phiếu", "Ngày BH", "Tên file", "Nội dung cần lưu ý", "Link phiếu"],
     [[i, p["tb"], p["so"], p["nb"], p["ten"], "; ".join(filter(None, [p.get("gc", ""), "Trích đoạn chưa đủ thông số - xem file gốc" if p.get("thieu") else ""])), f'https://drive.google.com/file/d/{p["id"]}/view'] for i, p in enumerate(ds_l, 1)],
     [6, 30, 24, 11, 34, 70, 40])

ws = wb.create_sheet("Tổng hợp", 0)
co_khoa = [t for t in D["thietBi"] if t["khoa"]]
k_co = {t["khoa"] for t in co_khoa if t["soPhieu"]}
k_tat = {t["khoa"] for t in co_khoa}
dem_khu = collections.Counter(khu(t["khoa"].split(":")[0]) for t in trung_ap)
dong = [
    ["Số phiếu chỉnh định đã trích (trang 1)", D["tongPhieu"]],
    ["Máy cắt / Recloser trên sơ đồ đọc được số hiệu", len(k_tat)],
    ["  - đã gắn được phiếu", len(k_co)],
    ["  - chưa có phiếu: trung áp", len(trung_ap)],
    ["       trong đó Thái Nguyên", dem_khu["Thái Nguyên"]],
    ["       trong đó Bắc Kạn", dem_khu["Bắc Kạn"]],
    ["  - chưa có phiếu: 110-220kV", len(cao_ap)],
    ["Phiếu chưa gắn được vào thiết bị trên sơ đồ", len(D["phieuKhongTB"])],
    ["Phiếu cần lưu ý (sai lệch / trích đoạn / bản lưu)", len(D["luuY"])],
]
bang(ws, "TỔNG HỢP ĐỐI CHIẾU PHIẾU CHỈNH ĐỊNH RƠ LE", ["Nội dung", "Số lượng"], dong, [58, 12])
ws.append([])
ws.append(["Ghi chú: thiết bị trong trạm nhận theo nhãn số máy cắt gần ký hiệu nhất; Recloser nhận theo nhãn \"MC <lộ><trạm>/<cột>\"."])
ws.append(["Không tính ký hiệu máy cắt dùng cho ngăn TU, cầu chì, tủ khách hàng (nhãn không phải số máy cắt)."])
wb.save(RA)
print(RA, "| trung áp", len(trung_ap), "| 110-220", len(cao_ap), "| phiếu chưa gắn", len(D["phieuKhongTB"]))
